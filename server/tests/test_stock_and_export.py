"""Suppliers, receiving with cost, write-offs, the Excel report and the shop logo."""

import io

from openpyxl import load_workbook
from PIL import Image

from tests.conftest import product_by_name
from tests.test_api import checkout


def png(width=600, height=300) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGBA", (width, height), (37, 99, 235, 255)).save(buffer, "PNG")
    return buffer.getvalue()


def test_suppliers(client, manager, cashier):
    created = client.post("/api/suppliers", headers=manager, json={"name": " Angkor Drinks Co ", "phone": "012 111 222"})
    assert created.status_code == 201 and created.json()["name"] == "Angkor Drinks Co"
    assert client.post("/api/suppliers", headers=manager, json={"name": "Angkor Drinks Co"}).status_code == 409
    supplier = created.json()
    updated = client.put(f"/api/suppliers/{supplier['id']}", headers=manager, json={"name": "Angkor Drinks Co", "is_active": False})
    assert updated.json()["is_active"] is False
    assert client.get("/api/suppliers", headers=cashier).status_code == 403


def test_receiving_with_cost_averages_the_product_cost(client, manager):
    supplier = client.post("/api/suppliers", headers=manager, json={"name": "Rice Mill"}).json()
    rice = product_by_name(client, manager, "Jasmine Rice 5kg")  # 15 on hand at $5.20
    body = {"product_id": rice["id"], "qty": 5, "unit_cost": 600, "supplier_id": supplier["id"], "note": "INV-88"}
    received = client.post("/api/inventory/stock-in", headers=manager, json=body).json()
    assert received["stock_qty"] == 20
    assert received["cost"] == (15 * 520 + 5 * 600) // 20  # $5.40

    move = client.get(f"/api/inventory/movements?product_id={rice['id']}", headers=manager).json()[0]
    assert move["kind"] == "stock_in" and move["unit_cost"] == 600 and move["supplier_name"] == "Rice Mill" and move["note"] == "INV-88"

    deliveries = client.get(f"/api/inventory/movements?supplier_id={supplier['id']}", headers=manager).json()
    assert [(m["product_name"], m["change"]) for m in deliveries] == [("Jasmine Rice 5kg", 5)]

    no_cost = client.post("/api/inventory/stock-in", headers=manager, json={"product_id": rice["id"], "qty": 1}).json()
    assert no_cost["cost"] == 540  # cost unchanged when no price is entered
    missing = client.post("/api/inventory/stock-in", headers=manager, json={**body, "supplier_id": 999})
    assert missing.status_code == 400

    report = client.get("/api/reports/summary", headers=manager).json()
    assert report["purchases"] == 5 * 600


def test_write_off_records_the_loss(client, manager):
    milk = product_by_name(client, manager, "Fresh Milk 1L")  # 14 at $1.60
    body = {"product_id": milk["id"], "qty": 2, "reason": "expired", "note": "Out of date 14 Oct"}
    after = client.post("/api/inventory/write-off", headers=manager, json=body).json()
    assert after["stock_qty"] == 12
    move = client.get(f"/api/inventory/movements?product_id={milk['id']}&kind=write_off", headers=manager).json()
    assert len(move) == 1 and move[0]["change"] == -2 and move[0]["note"] == "Expired: Out of date 14 Oct"

    too_many = client.post("/api/inventory/write-off", headers=manager, json={**body, "qty": 99})
    assert too_many.status_code == 400
    bad_reason = client.post("/api/inventory/write-off", headers=manager, json={**body, "reason": "stolen?"})
    assert bad_reason.status_code == 422

    report = client.get("/api/reports/summary", headers=manager).json()
    assert report["stock_losses"] == 2 * 160
    assert report["stock_value"] > 0


def test_excel_report(client, manager):
    assert client.post("/api/shifts/open", headers=manager, json={"opening_cash": 0}).status_code == 201
    rice = product_by_name(client, manager, "Jasmine Rice 5kg")
    checkout(client, manager, [{"product_id": rice["id"], "qty": 2}], [{"method": "cash", "amount": 2000}])
    milk = product_by_name(client, manager, "Fresh Milk 1L")
    client.post("/api/inventory/write-off", headers=manager, json={"product_id": milk["id"], "qty": 1, "reason": "damaged"})

    response = client.get("/api/reports/export", headers=manager)
    assert response.status_code == 200
    assert "attachment; filename=sales-" in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Summary", "By day", "Products", "Categories", "Payments", "Cashiers", "Busy hours", "Receipts", "Stock losses"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=4, values_only=True)}
    assert summary["Net sales"] == 14.30 and summary["Sales (receipts)"] == 1
    assert summary["Stock losses (damaged, expired, lost)"] == 1.60
    receipts = list(wb["Receipts"].iter_rows(min_row=2, values_only=True))
    assert len(receipts) == 1 and receipts[0][8] == 14.30 and receipts[0][10] == "Cash"
    assert list(wb["Stock losses"].iter_rows(min_row=2, values_only=True))[0][1:4] == ("Fresh Milk 1L", 1, "Damaged")
    assert client.get("/api/reports/export", headers=manager, params={"date_from": "2026-01-02", "date_to": "2026-01-01"}).status_code == 400


def test_logo_upload_and_remove(client, admin, manager, monkeypatch, tmp_path):
    from app import images
    monkeypatch.setattr(images, "LOGO_DIR", tmp_path)
    assert client.post("/api/settings/logo", headers=manager, files={"file": ("logo.png", png(), "image/png")}).status_code == 403

    settings = client.post("/api/settings/logo", headers=admin, files={"file": ("logo.png", png(), "image/png")}).json()
    assert settings["logo_url"].startswith("/uploads/logo/logo-")
    stored = tmp_path / settings["logo_url"].rsplit("/", 1)[1]
    assert Image.open(stored).size == (512, 256)  # shrunk, shape kept

    bad = client.post("/api/settings/logo", headers=admin, files={"file": ("logo.png", b"not an image", "image/png")})
    assert bad.status_code == 400
    assert client.get("/api/settings", headers=manager).json()["logo_url"] == settings["logo_url"]

    removed = client.delete("/api/settings/logo", headers=admin).json()
    assert removed["logo_url"] == "" and not stored.exists()
