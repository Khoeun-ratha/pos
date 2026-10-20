import io
from datetime import date

from openpyxl import load_workbook

from conftest import login

TODAY = date.today().isoformat()


def expense(client, headers, amount, category="Electricity", **extra):
    return client.post("/api/expenses", json={"spent_on": TODAY, "category": category, "amount": amount, **extra}, headers=headers)


def test_expenses_lower_the_profit_and_are_listed_by_category(client, manager, cashier):
    assert expense(client, cashier, 500).status_code == 403  # cashiers do not record expenses by default
    assert expense(client, manager, 4500, note="October bill").status_code == 201
    assert expense(client, manager, 120000, category="Rent").status_code == 201
    assert expense(client, manager, 0).status_code == 422

    summary = client.get("/api/reports/summary", headers=manager).json()
    assert summary["expenses"] == 124500
    assert summary["net_profit"] == summary["profit"] - 124500

    report = client.get("/api/reports/expenses", headers=manager).json()
    assert [(c["name"], c["amount"]) for c in report["categories"]] == [("Rent", 120000), ("Electricity", 4500)]
    assert len(report["items"]) == 2 and report["items"][0]["recorded_by"] == "Store Manager"

    categories = client.get("/api/expenses/categories", headers=manager).json()
    assert categories[0] == "Rent" and "Salaries" in categories
    assert expense(client, manager, 300, category="Security guard").status_code == 201
    assert "Security guard" in client.get("/api/expenses/categories", headers=manager).json()

    xlsx = client.get("/api/reports/export", headers=manager)
    sheet = load_workbook(io.BytesIO(xlsx.content))["Expenses"]
    assert sheet.max_row == 4 and sheet["B2"].value == "Electricity"


def test_expense_paid_from_the_drawer_follows_the_shift(client, manager):
    assert expense(client, manager, 800, category="Ice & supplies", from_drawer=True).status_code == 400  # no shift
    client.post("/api/shifts/open", json={"opening_cash": 5000}, headers=manager)
    created = expense(client, manager, 800, category="Ice & supplies", from_drawer=True).json()
    assert created["from_drawer"]
    shift = client.get("/api/shifts/current", headers=manager).json()
    assert shift["cash_out"] == 800 and shift["expected_cash"] == 4200

    # while the shift is open the amount can change, and the drawer follows
    body = {"spent_on": TODAY, "category": "Ice & supplies", "amount": 1000, "note": "2 bags"}
    assert client.put(f"/api/expenses/{created['id']}", json=body, headers=manager).status_code == 200
    assert client.get("/api/shifts/current", headers=manager).json()["expected_cash"] == 4000

    # deleting it puts the cash back
    other = expense(client, manager, 250, category="Transport", from_drawer=True).json()
    assert client.delete(f"/api/expenses/{other['id']}", headers=manager).status_code == 204
    assert client.get("/api/shifts/current", headers=manager).json()["expected_cash"] == 4000

    # once the shift is closed, the amount is fixed (the note can still be corrected)
    client.post("/api/shifts/current/close", json={"counted_usd": 4000}, headers=manager)
    assert client.put(f"/api/expenses/{created['id']}", json={**body, "amount": 1}, headers=manager).status_code == 400
    assert client.put(f"/api/expenses/{created['id']}", json={**body, "note": "3 bags"}, headers=manager).status_code == 200
    assert client.delete(f"/api/expenses/{created['id']}", headers=manager).status_code == 400
    closed = client.get("/api/shifts", headers=manager).json()[0]
    assert closed["variance"] == 0 and closed["cash_out"] == 1000


def test_manager_manages_cashiers_only(client, manager, admin, cashier):
    users = client.get("/api/users", headers=manager)
    assert users.status_code == 200
    by_name = {u["username"]: u for u in users.json()}

    new = client.post("/api/users", json={"username": "dara", "full_name": "Dara", "password": "dara1234", "pin": "5678", "role": "cashier"},
                      headers=manager)
    assert new.status_code == 201, new.text
    assert client.post("/api/auth/login", json={"username": "dara", "password": "dara1234"}).status_code == 200
    # reset password and PIN, deactivate
    dara = new.json()["id"]
    assert client.patch(f"/api/users/{dara}", json={"password": "new12345", "pin": "8765"}, headers=manager).status_code == 200
    assert client.post("/api/auth/login", json={"username": "dara", "password": "new12345"}).status_code == 200
    assert client.patch(f"/api/users/{dara}", json={"is_active": False}, headers=manager).status_code == 200
    assert client.post("/api/auth/login", json={"username": "dara", "password": "new12345"}).status_code in (401, 403)

    # not managers, admins, or promotions
    assert client.post("/api/users", json={"username": "boss", "password": "boss1234", "role": "manager"}, headers=manager).status_code == 403
    assert client.patch(f"/api/users/{by_name['admin']['id']}", json={"password": "hacked123"}, headers=manager).status_code == 403
    assert client.patch(f"/api/users/{by_name['cashier2']['id']}", json={"role": "manager"}, headers=manager).status_code == 403
    assert client.patch(f"/api/users/{by_name['manager']['id']}", json={"full_name": "Me"}, headers=manager).status_code == 403
    assert client.post("/api/auth/login", json={"username": "admin", "password": "admin123"}).status_code == 200

    # cashiers do not manage staff; the admin can take the right away from managers
    assert client.get("/api/users", headers=cashier).status_code == 403
    client.put("/api/permissions", json={"manager": ["reports"], "cashier": []}, headers=admin)
    assert client.get("/api/users", headers=login(client, "manager")).status_code == 403
    # the admin still manages everyone
    assert client.patch(f"/api/users/{by_name['manager']['id']}", json={"full_name": "Head of shop"}, headers=admin).status_code == 200
