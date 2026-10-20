"""Sales report as an Excel workbook: a summary sheet plus one sheet per breakdown, every receipt and stock losses."""

import io
from collections.abc import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Expense, Sale, StockMovement
from .schemas import ReportSummary

USD = '"$"#,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
METHOD_NAMES = {"cash": "Cash", "card": "Card", "khqr": "KHQR"}


def _dollars(cents: int | None) -> float | None:
    return None if cents is None else cents / 100


def _table(ws: Worksheet, headers: Sequence[str], rows: Iterable[Sequence], money: Sequence[int] = (), start_row: int = 1) -> None:
    """Header row, data rows, money columns (0-based) formatted as dollars, sensible widths, frozen header."""
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    count = 0
    for r, row in enumerate(rows, start=start_row + 1):
        count += 1
        for c, value in enumerate(row):
            cell = ws.cell(row=r, column=c + 1, value=_dollars(value) if c in money else value)
            if c in money:
                cell.number_format = USD
    for col, title in enumerate(headers, start=1):
        values = [str(ws.cell(row=r, column=col).value or "") for r in range(start_row, start_row + count + 1)]
        ws.column_dimensions[get_column_letter(col)].width = min(max(len(v) for v in values) + 3, 50)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def report_workbook(report: ReportSummary, sales: Sequence[Sale], losses: Sequence[StockMovement], store_name: str,
                    expenses: Sequence[Expense] = ()) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{store_name}: sales report"
    ws["A1"].font = Font(bold=True, size=14)
    period = str(report.date_from) if report.date_from == report.date_to else f"{report.date_from} to {report.date_to}"
    ws["A2"] = period
    summary = [
        ("Sales (receipts)", report.sales_count, False),
        ("Gross sales", report.gross_sales, True),
        ("Refunds", report.refunds, True),
        ("Net sales", report.net_sales, True),
        ("VAT collected", report.tax, True),
        ("Discounts given", report.discounts, True),
        ("Cost of goods sold", report.cost, True),
        ("Gross profit", report.profit, True),
        ("Expenses (rent, salaries...)", report.expenses, True),
        ("Profit after expenses", report.net_profit, True),
        ("Average sale", report.average_sale, True),
        ("Stock received (at cost)", report.purchases, True),
        ("Stock losses (damaged, expired, lost)", report.stock_losses, True),
        ("Stock value now (at cost)", report.stock_value, True),
    ]
    for i, (label, value, money) in enumerate(summary, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=label in ("Net sales", "Gross profit", "Profit after expenses"))
        cell = ws.cell(row=i, column=2, value=_dollars(value) if money else value)
        if money:
            cell.number_format = USD
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 16

    _table(wb.create_sheet("By day"), ["Day", "Sales", "Net amount"],
           ((d.day, d.count, d.amount) for d in report.by_day), money=[2])
    _table(wb.create_sheet("Products"), ["Product", "Qty", "Amount"],
           ((p.name, p.qty, p.amount) for p in report.top_products), money=[2])
    _table(wb.create_sheet("Categories"), ["Category", "Qty", "Amount"],
           ((c.name, c.qty, c.amount) for c in report.by_category), money=[2])
    _table(wb.create_sheet("Payments"), ["Method", "Payments", "Amount"],
           ((p.name, p.count, p.amount) for p in report.by_payment), money=[2])
    _table(wb.create_sheet("Cashiers"), ["Cashier", "Sales", "Amount"],
           ((c.name, c.count, c.amount) for c in report.by_cashier), money=[2])
    _table(wb.create_sheet("Busy hours"), ["Hour", "Sales", "Amount"],
           ((f"{h.hour:02d}:00", h.count, h.amount) for h in report.by_hour if h.count or h.amount), money=[2])

    _table(
        wb.create_sheet("Receipts"),
        ["Receipt", "Date", "Cashier", "Customer", "Items", "Subtotal", "Discount", "VAT", "Total", "Refunded", "Paid by", "Status"],
        (
            (s.receipt_no, s.created_at, s.cashier_name, s.customer_name or "", sum(i.qty for i in s.items),
             s.subtotal, s.discount_total, s.tax_total, s.total, s.refunded_total,
             ", ".join(sorted({METHOD_NAMES.get(p.method, p.method) for p in s.payments})), s.status.replace("_", " "))
            for s in sales
        ),
        money=[5, 6, 7, 8, 9],
    )
    for row in wb["Receipts"].iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd hh:mm"

    _table(
        wb.create_sheet("Stock losses"),
        ["Date", "Product", "Qty", "Reason", "Unit cost", "Total cost"],
        ((m.created_at, m.product_name, -m.change, m.note or "", m.unit_cost, (m.unit_cost or 0) * -m.change) for m in losses),
        money=[4, 5],
    )
    for row in wb["Stock losses"].iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "yyyy-mm-dd hh:mm"

    _table(
        wb.create_sheet("Expenses"),
        ["Date", "Category", "Note", "Amount", "From the cash drawer", "Recorded by"],
        ((e.spent_on, e.category, e.note or "", e.amount, "yes" if e.from_drawer else "", e.recorded_by) for e in expenses),
        money=[3],
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
